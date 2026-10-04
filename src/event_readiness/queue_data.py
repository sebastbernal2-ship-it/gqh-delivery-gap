"""Normalize LBNL annual queue files without inventing publication clocks or stable identities."""
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from .acquisition import sha
from .pull_candidates import record


def value_text(value):
    if value is None:
        return None
    s = str(value).strip()
    return None if s.lower() in {"", "na", "n/a", "none", "nan", "not assigned"} else s


def date_value(value, epoch, year_hint=None):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # Styled Excel dates arrive as datetime. Unformatted numbers may use a
        # different epoch or represent a year. Decode only with a companion year
        # that independently matches the workbook's Excel epoch conversion.
        from openpyxl.utils.datetime import from_excel
        try:
            year = int(year_hint)
            if 1900 <= year <= 2100 and 20000 <= value < 73050:
                parsed = from_excel(value, epoch=epoch).date()
                if parsed.year == year:
                    return parsed.isoformat()
        except (TypeError, ValueError, OverflowError):
            pass
        return None
    text = value_text(value)
    if text:
        try:
            return datetime.fromisoformat(text).date().isoformat()
        except ValueError:
            pass
        for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%m/%d/%y", "%Y-%m-%d %H:%M:%S"):
            try:
                return datetime.strptime(text, fmt).date().isoformat()
            except ValueError:
                pass
    return None


def scalar(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def queue_records(path: Path, year: int, receipt: dict):
    from openpyxl import load_workbook
    if sha(path.read_bytes()) != receipt["sha256"]:
        raise ValueError("queue file hash differs from receipt")
    with path.open("rb") as stream:
        wb = load_workbook(stream, read_only=True, data_only=True)
        candidates = [s for s in wb.worksheets if s.title.lower() == "data" or s.title == "03. Complete Queue Data"]
        if year == 2020:
            candidates = [wb[name] for name in ("active", "withdrawn", "completed")]
        elif len(candidates) != 1:
            raise ValueError("ambiguous queue data sheet")
        for sheet in candidates:
            yield from sheet_records(sheet, year, receipt, wb.epoch)
        wb.close()


def sheet_records(sheet, year, receipt, epoch):
    rows = sheet.iter_rows(values_only=True)
    header = None
    for offset in range(3):
        candidate = next(rows)
        if candidate[0] == "q_id":
            header = [str(c).lower() for c in candidate]
            break
    if not header or len(header) != len(set(header)):
        raise ValueError("invalid queue header")
    required = {"q_id", "entity", "q_status", "state", "type_clean"}
    if not required <= set(header):
        raise ValueError("queue schema changed")
    if not {"q_date", "q_date_clean"} & set(header) or not {"prop_date", "proposed_on_date"} & set(header):
        raise ValueError("queue date schema changed")
    for number, cells in enumerate(rows, offset + 2):
        if all(v is None for v in cells):
            continue
        raw = dict(zip(header, map(scalar, cells)))
        operator, project = value_text(raw["entity"]), value_text(raw["q_id"])
        canonical_fields = {**raw, "q_date": raw.get("q_date", raw.get("q_date_clean")),
                            "prop_date": raw.get("prop_date", raw.get("proposed_on_date"))}
        hints = {"q_date": raw.get("q_year"), "prop_date": raw.get("prop_year", raw.get("proposed_on_year"))}
        dates = {k: date_value(canonical_fields.get(k), epoch, hints.get(k)) for k in ("q_date", "prop_date", "on_date", "wd_date", "ia_date")}
        flags = ["historical_publication_unverified", "annual_snapshot_not_monthly_history"]
        for key in dates:
            if dates[key] and isinstance(canonical_fields.get(key), (int, float)):
                flags.append(key + "_excel_epoch_matches_companion_year")
            if value_text(canonical_fields.get(key)) is not None and dates[key] is None:
                flags.append(key + "_unparsed")
        if not operator or not project:
            flags.append("missing_candidate_project_identity")
        normalized = {"operator": operator, "project_id": project,
            "candidate_project_key": operator + ":" + project if operator and project else None,
            "status_raw": raw.get("q_status"), "state": (value_text(raw.get("state")) or "").upper() or None,
            "fuel_raw": raw.get("type_clean"), "point_of_interconnection": raw.get("poi_name"),
            "county_raw": raw.get("county", raw.get("county_1")), "cluster_raw": raw.get("cluster"),
            "submission_date": dates["q_date"], "proposed_service_date": dates["prop_date"],
            "operation_date": dates["on_date"], "withdrawal_date": dates["wd_date"],
            "interconnection_agreement_date": dates["ia_date"],
            "submission_year_raw": raw.get("q_year"), "proposed_year_raw": raw.get("prop_year", raw.get("proposed_on_year")),
            "mw_components_raw": [raw.get(f"mw_{i}", raw.get(f"mw{i}")) for i in (1, 2, 3)]}
        yield record("lbnl_queue_project", {"snapshot_as_of": f"{year}-12-31", "sheet": sheet.title,
            "sheet_row": number, "historical_availability_verified": False}, receipt,
            raw, normalized=normalized, quality_flags=flags)


def summarize(records):
    rows = list(records)
    keys = Counter(r["normalized"]["candidate_project_key"] for r in rows if r["normalized"]["candidate_project_key"])
    return {"rows": len(rows), "distinct_states": sorted({r["normalized"]["state"] for r in rows if r["normalized"]["state"]}),
        "distinct_fuels": sorted({str(r["normalized"]["fuel_raw"]) for r in rows if value_text(r["normalized"]["fuel_raw"])}),
        "usable_submission_dates": sum(bool(r["normalized"]["submission_date"]) for r in rows),
        "usable_proposed_service_dates": sum(bool(r["normalized"]["proposed_service_date"]) for r in rows),
        "duplicate_candidate_identity_groups": sum(n > 1 for n in keys.values()),
        "rows_missing_candidate_identity": sum(not r["normalized"]["candidate_project_key"] for r in rows),
        "verified_publication_dates": 0, "point_in_time_ready": False}
