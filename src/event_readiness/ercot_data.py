"""Extract dated ERCOT project observations, preserving headers, cells and document versions."""
import argparse
import calendar
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
import re
import shutil

from .acquisition import atomic_json, records_file, sha, strict_json
from .pull_candidates import record
from .queue_data import date_value, scalar, value_text


def report_month(name):
    months = {calendar.month_name[i][:3].lower(): i for i in range(1, 13)}
    matches = re.findall(r"([A-Za-z]+)[_ ]*(20\d{2})", name)
    dates = [f"{year}-{months[month[:3].lower()]:02}" for month, year in matches if month[:3].lower() in months]
    if len(dates) != 1:
        raise ValueError("ambiguous ERCOT report month")
    return dates[0]


def milestone(value, epoch):
    parsed = date_value(value, epoch)
    # Explicit missing-date marker described in ERCOT workbook notes.
    return None if parsed and parsed <= "1900-01-02" else parsed


def extract_sheet(sheet, workbook_record, receipt, epoch):
    headers = None
    for number, cells in enumerate(sheet.iter_rows(values_only=True), 1):
        if headers is None:
            labels = [value_text(v) for v in cells]
            if "INR" in labels:
                headers = labels
                id_column = labels.index("INR")
            elif number >= 60:
                raise ValueError("project sheet has no recognized INR header")
            continue
        identity = value_text(cells[id_column])
        if not identity or identity == "INR":
            continue
        raw = {"headers": headers, "cells": [scalar(v) for v in cells]}
        named = {h: scalar(cells[i]) for i, h in enumerate(headers) if h}
        flags = ["original_publication_unverified", "limited_public_queue_coverage", "project_split_links_unresolved"]
        if not re.fullmatch(r"\d{2}INR[\w.-]+", identity, re.IGNORECASE):
            flags.append("unrecognized_project_id_format")
        projected = milestone(named.get("Projected COD"), epoch)
        if value_text(named.get("Projected COD")) is not None and projected is None:
            flags.append("projected_cod_missing_or_unparsed")
        normalized = {
            "operator": "ERCOT", "project_id": identity, "candidate_project_key": "ERCOT:" + identity,
            "state": "TX", "project_name": named.get("Project Name"), "county_raw": named.get("County"),
            "point_of_interconnection": named.get("POI Location"), "fuel_raw": named.get("Fuel"),
            "technology_raw": named.get("Technology"), "capacity_mw_raw": named.get("Capacity (MW)"),
            "interconnecting_entity_raw": named.get("Interconnecting Entity"),
            "phase_raw": named.get("GIM Study Phase", named.get("GINR Study Phase")),
            "proposed_service_date": projected, "submission_date": None,
            "fis_requested_date": milestone(named.get("FIS Requested"), epoch),
            "interconnection_agreement_date": milestone(named.get("IA Signed"), epoch),
            "approved_energization_date": milestone(named.get("Approved for Energization"), epoch),
            "approved_synchronization_date": milestone(named.get("Approved for Synchronization"), epoch),
        }
        context = {**workbook_record["context"], "report_month": report_month(workbook_record["data"]["FriendlyName"]),
                   "sheet": sheet.title, "sheet_row": number}
        yield record("ercot_queue_project", context, receipt, raw, normalized=normalized,
                     reported_posted_at_utc=workbook_record["reported_posted_at_utc"], quality_flags=flags)
    if headers is None:
        raise ValueError("missing INR header")


def normalize(source, output):
    from openpyxl import load_workbook
    if (output / "complete.json").exists():
        raise ValueError("completed output exists")
    complete = strict_json((source / "complete.json").read_bytes())
    if sha((source / "records.jsonl").read_bytes()) != complete["sha256"]:
        raise ValueError("workbook inventory checksum mismatch")
    for folder in ("objects", "requests"):
        shutil.copytree(source / folder, output / folder, dirs_exist_ok=True)
    documents = [strict_json(line) for line in (source / "records.jsonl").read_bytes().splitlines()]
    receipts = {r["sha256"]: r for r in (strict_json(p.read_bytes()) for p in (output / "requests").glob("*.json"))}
    quality = {"workbooks": len(documents), "sheets": [], "submission_dates_available": False}
    counts = Counter()
    def rows():
        for i, document in enumerate(documents, 1):
            path = output / "objects" / document["source_sha256"]
            if sha(path.read_bytes()) != document["source_sha256"]:
                raise ValueError("workbook hash mismatch")
            receipt = receipts[document["source_sha256"]]
            with path.open("rb") as f:
                wb = load_workbook(f, read_only=True, data_only=True)
                sheets = [s for s in wb if s.title.startswith("Project Details")]
                if not sheets:
                    raise ValueError("no project detail sheet")
                for sheet in sheets:
                    n = 0
                    for row in extract_sheet(sheet, document, receipt, wb.epoch):
                        n += 1
                        counts[row["context"]["report_month"]] += 1
                        yield row
                    quality["sheets"].append({"doc_id": document["context"]["doc_id"], "sheet": sheet.title, "rows": n})
                wb.close()
            if i % 20 == 0:
                print(f"normalized ERCOT {i}/{len(documents)}", flush=True)
    result = records_file(output / "records.jsonl", rows())
    quality.update(report_month_counts=dict(sorted(counts.items())),
                   scope="Project Details sheets only; inactive/cancellation/commissioning sheets retained in raw workbooks")
    result.update(kind="ercot_project_details", strategy_ready=False, quality=quality,
                  parent_inventory_sha256=complete["sha256"], completed_at_utc=datetime.now(timezone.utc).isoformat())
    atomic_json(output / "complete.json", result)
    print(f"ERCOT normalized: {result['rows']} rows", flush=True)


if __name__ == "__main__":
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("source", type=Path); p.add_argument("output", type=Path)
    args = p.parse_args()
    normalize(args.source, args.output)
